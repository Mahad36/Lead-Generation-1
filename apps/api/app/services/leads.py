import csv
import io
from datetime import UTC, datetime
from typing import Any

from fastapi import HTTPException, status
from sqlalchemy import and_, func, or_
from sqlalchemy.orm import Session, joinedload

from app.logging_config import get_logger
from app.models import Lead, LeadEmail, LeadImport, LeadPhone, Website, WebsiteAnalysis
from app.schemas.leads import LeadCreate, LeadUpdate
from app.services.audit import write_audit
from app.services.duplicates import find_duplicate_lead
from app.services.lead_outreach import outreach_for_lead
from app.services.scoring import calculate_lead_score
from app.utils.normalize import (
    extract_domain,
    normalize_address,
    normalize_business_name,
    normalize_email,
    normalize_phone,
    normalize_url,
)

logger = get_logger("leads")


def outreach_status(email: str | None) -> str:
    return "ready" if (email or "").strip() else "no_email"

EXPORT_COLUMNS = [
    "business_name",
    "category",
    "country",
    "city",
    "address",
    "phone",
    "email",
    "website",
    "platform",
    "website_score",
    "lead_score",
    "email_status",
]


def purge_leads_without_email(db: Session, user_id: str) -> int:
    rows = (
        db.query(Lead)
        .filter(
            Lead.user_id == user_id,
            Lead.status != "processing",
            Lead.scraping_job_id.is_(None),
            or_(Lead.primary_email == "", Lead.primary_email.is_(None), Lead.status == "no_email"),
        )
        .all()
    )
    for lead in rows:
        db.delete(lead)
    if rows:
        db.commit()
    return len(rows)


def discard_lead_without_email(db: Session, lead: Lead) -> bool:
    if (lead.primary_email or "").strip():
        return False
    if lead.scraping_job_id:
        lead.status = "no_email"
        lead.email_status = "none"
        db.add(lead)
        db.commit()
        logger.info("lead_no_email_kept", lead_id=lead.id)
        return False
    batch_id = lead.import_batch_id
    db.delete(lead)
    db.commit()
    from app.services.lead_pipeline import refresh_import_batch

    refresh_import_batch(db, batch_id)
    logger.info("lead_discarded_no_email", lead_id=lead.id)
    return True


def import_batch_filter(batch: LeadImport):
    """Leads of an uploaded file: the ones it created plus existing leads its rows matched (linked)."""
    linked = list(batch.linked_lead_ids or [])
    return or_(Lead.import_batch_id == batch.id, Lead.id.in_(linked)) if linked else Lead.import_batch_id == batch.id


def _apply_filters(query, user_id: str, params: dict[str, Any]):
    query = query.filter(Lead.user_id == user_id)
    if not params.get("scraping_job_id"):
        query = query.filter(Lead.primary_email != "", Lead.status != "no_email")
    if q := (params.get("q") or "").strip():
        like = f"%{q}%"
        query = query.filter(
            or_(
                Lead.business_name.ilike(like),
                Lead.primary_email.ilike(like),
                Lead.website_domain.ilike(like),
                Lead.category.ilike(like),
                Lead.city.ilike(like),
            )
        )
    for field in (
        "country",
        "city",
        "category",
        "platform",
        "email_status",
        "email_verification_status",
        "status",
        "source",
        "scraping_job_id",
    ):
        value = params.get(field)
        if value:
            query = query.filter(getattr(Lead, field) == value)
    if params.get("import_batch_id"):
        batch = query.session.get(LeadImport, params["import_batch_id"])
        query = query.filter(import_batch_filter(batch) if batch else Lead.import_batch_id == params["import_batch_id"])
    if params.get("has_email") == "false":
        query = query.filter(Lead.id.is_(None))
    if params.get("has_website") == "true":
        query = query.filter(Lead.website_url != "")
    elif params.get("has_website") == "false":
        query = query.filter(or_(Lead.website_url == "", Lead.website_url.is_(None)))
    if not params.get("import_batch_id") and not params.get("scraping_job_id") and not params.get("include_files"):
        if params.get("source") != "manual":
            query = query.filter(Lead.import_batch_id.is_(None), Lead.scraping_job_id.is_(None))
    if params.get("min_lead_score") is not None:
        query = query.filter(Lead.lead_score >= int(params["min_lead_score"]))
    if params.get("max_lead_score") is not None:
        query = query.filter(Lead.lead_score <= int(params["max_lead_score"]))
    if params.get("min_website_score") is not None:
        query = query.filter(Lead.website_score >= int(params["min_website_score"]))
    if params.get("max_website_score") is not None:
        query = query.filter(Lead.website_score <= int(params["max_website_score"]))
    return query


SORT_MAP = {
    "business_name": Lead.business_name,
    "category": Lead.category,
    "city": Lead.city,
    "country": Lead.country,
    "platform": Lead.platform,
    "website_score": Lead.website_score,
    "lead_score": Lead.lead_score,
    "scraped_at": Lead.scraped_at,
    "created_at": Lead.created_at,
    "email_status": Lead.email_status,
}


def list_leads(db: Session, user_id: str, params: dict[str, Any]) -> tuple[list[Lead], int]:
    purge_leads_without_email(db, user_id)
    page = max(int(params.get("page") or 1), 1)
    page_size = min(max(int(params.get("page_size") or 25), 1), 200)
    query = _apply_filters(db.query(Lead), user_id, params)
    total = query.with_entities(func.count(Lead.id)).scalar() or 0
    sort = params.get("sort") or "created_at"
    order = (params.get("order") or "desc").lower()
    column = SORT_MAP.get(sort, Lead.created_at)
    query = query.order_by(column.asc() if order == "asc" else column.desc())
    items = query.offset((page - 1) * page_size).limit(page_size).all()
    return items, total


def get_lead(db: Session, user_id: str, lead_id: str) -> Lead:
    lead = (
        db.query(Lead)
        .options(joinedload(Lead.emails), joinedload(Lead.phones), joinedload(Lead.website))
        .filter(Lead.id == lead_id, Lead.user_id == user_id)
        .first()
    )
    if not lead:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Lead not found")
    return lead


def create_lead(db: Session, user_id: str, data: LeadCreate, meta: dict) -> Lead:
    duplicate = find_duplicate_lead(
        db,
        user_id,
        website_url=data.website_url,
        phone=data.primary_phone,
        business_name=data.business_name,
        address=data.address,
        email=data.primary_email,
        google_place_id=data.google_place_id,
    )
    if duplicate:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Duplicate lead")
    if not (data.business_name or "").strip() or not (data.primary_email or "").strip():
        raise HTTPException(status_code=400, detail="Business name and email are required")
    lead = persist_lead(db, user_id, data.model_dump(), source="manual")
    if not lead:
        raise HTTPException(status_code=400, detail="Lead was skipped because it has no email address")
    write_audit(db, user_id=user_id, action="lead.created", entity_type="lead", entity_id=lead.id, **meta)
    db.commit()
    db.refresh(lead)
    if lead.website_url:
        from app.services.lead_pipeline import enqueue_lead_website_processing

        enqueue_lead_website_processing([lead.id])
    return lead


def persist_lead(
    db: Session,
    user_id: str,
    data: dict,
    *,
    source: str = "google_maps",
    scraping_job_id: str | None = None,
    import_batch_id: str | None = None,
    allow_without_email: bool = False,
) -> Lead | None:
    phone = normalize_phone(data.get("primary_phone") or data.get("phone") or "")
    email = normalize_email(data.get("primary_email") or data.get("email") or "")
    website = normalize_url(data.get("website_url") or data.get("website") or "")
    domain = extract_domain(website)
    name = data.get("business_name") or ""
    if not name or (not email and not allow_without_email):
        return None
    if data.get("status"):
        lead_status = data["status"]
    elif website and not email:
        lead_status = "processing"
    else:
        lead_status = outreach_status(email)
    lead = Lead(
        user_id=user_id,
        scraping_job_id=scraping_job_id,
        import_batch_id=import_batch_id,
        business_name=name,
        contact_name=(data.get("contact_name") or "").strip()[:255],
        normalized_name=normalize_business_name(name),
        category=data.get("category") or "",
        country=data.get("country") or "",
        city=data.get("city") or "",
        address=data.get("address") or "",
        normalized_address=normalize_address(data.get("address") or ""),
        google_maps_url=data.get("google_maps_url") or "",
        google_place_id=data.get("google_place_id") or "",
        rating=data.get("rating"),
        review_count=data.get("review_count"),
        website_url=website,
        website_domain=domain,
        primary_phone=phone,
        primary_email=email,
        email_status="found" if email else "none",
        notes=data.get("notes") or "",
        source=source,
        source_search=data.get("source_search") or "",
        status=lead_status,
        scraped_at=datetime.now(UTC) if source == "google_maps" else None,
        extra=data.get("extra") or {},
    )
    db.add(lead)
    db.flush()
    if email:
        db.add(LeadEmail(lead_id=lead.id, email=email, is_primary=True, source_url=website))
    if phone:
        db.add(LeadPhone(lead_id=lead.id, phone=phone, is_primary=True, source=source))
    score, reasons = calculate_lead_score(lead)
    lead.lead_score = score
    lead.lead_score_reasons = reasons
    return lead


def update_lead(db: Session, user_id: str, lead_id: str, data: LeadUpdate) -> Lead:
    lead = get_lead(db, user_id, lead_id)
    updates = data.model_dump(exclude_unset=True)
    if "primary_phone" in updates and updates["primary_phone"] is not None:
        updates["primary_phone"] = normalize_phone(updates["primary_phone"])
    if "primary_email" in updates and updates["primary_email"] is not None:
        updates["primary_email"] = normalize_email(updates["primary_email"])
        lead.email_status = "found" if updates["primary_email"] else "none"
    if "website_url" in updates and updates["website_url"] is not None:
        updates["website_url"] = normalize_url(updates["website_url"])
        lead.website_domain = extract_domain(updates["website_url"])
    if "business_name" in updates and updates["business_name"]:
        lead.normalized_name = normalize_business_name(updates["business_name"])
    if "address" in updates and updates["address"] is not None:
        lead.normalized_address = normalize_address(updates["address"])
    for key, value in updates.items():
        setattr(lead, key, value)
    score, reasons = calculate_lead_score(lead)
    lead.lead_score = score
    lead.lead_score_reasons = reasons
    db.commit()
    db.refresh(lead)
    return lead


def delete_leads(db: Session, user_id: str, ids: list[str], meta: dict) -> int:
    q = db.query(Lead).filter(Lead.user_id == user_id, Lead.id.in_(ids))
    count = q.count()
    q.delete(synchronize_session=False)
    write_audit(
        db,
        user_id=user_id,
        action="leads.deleted",
        entity_type="lead",
        details={"count": count, "ids": ids},
        **meta,
    )
    db.commit()
    return count


def export_leads_csv(db: Session, user_id: str, params: dict[str, Any], ids: list[str] | None) -> str:
    query = _apply_filters(db.query(Lead), user_id, params)
    if ids:
        query = query.filter(Lead.id.in_(ids))
    leads = query.order_by(Lead.created_at.desc()).limit(10000).all()
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=EXPORT_COLUMNS)
    writer.writeheader()
    for lead in leads:
        writer.writerow(
            {
                "business_name": lead.business_name,
                "category": lead.category,
                "country": lead.country,
                "city": lead.city,
                "address": lead.address,
                "phone": lead.primary_phone,
                "email": lead.primary_email,
                "website": lead.website_url,
                "platform": lead.platform,
                "website_score": lead.website_score if lead.website_score is not None else "",
                "lead_score": lead.lead_score if lead.lead_score is not None else "",
                "email_status": lead.email_status,
            }
        )
    return buffer.getvalue()


HEADER_ALIASES = {
    "business name": "business_name",
    "business_name": "business_name",
    "name": "business_name",
    "contact name": "contact_name",
    "contact_name": "contact_name",
    "contact": "contact_name",
    "contact person": "contact_name",
    "owner": "contact_name",
    "owner name": "contact_name",
    "first name": "contact_name",
    "category": "category",
    "country": "country",
    "city": "city",
    "address": "address",
    "phone": "primary_phone",
    "primary_phone": "primary_phone",
    "email": "primary_email",
    "primary_email": "primary_email",
    "website": "website_url",
    "website_url": "website_url",
    "google_maps_url": "google_maps_url",
    "notes": "notes",
}


def import_leads_file(db: Session, user_id: str, raw: bytes, meta: dict, filename: str = "leads.csv") -> dict:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xls"):
        content = _excel_to_csv(raw)
    else:
        content = raw.decode("utf-8-sig")
    return import_leads_csv(db, user_id, content, meta, filename=filename)


def _excel_to_csv(raw: bytes) -> str:
    import io

    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Could not read this spreadsheet. Upload an .xlsx file or save the Excel file as CSV.",
        ) from exc
    sheet = book.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="Spreadsheet is empty")
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    for row in rows:
        writer.writerow(["" if cell is None else str(cell) for cell in row])
    return buffer.getvalue()


def import_leads_csv(
    db: Session, user_id: str, content: str, meta: dict, filename: str = "leads.csv"
) -> dict:
    reader = csv.DictReader(io.StringIO(content))
    if not reader.fieldnames:
        raise HTTPException(status_code=400, detail="CSV is missing a header row")
    mapped_headers = {h: HEADER_ALIASES.get((h or "").strip().lower(), "") for h in reader.fieldnames}
    if "business_name" not in mapped_headers.values():
        raise HTTPException(
            status_code=400,
            detail="CSV must include a business name column (business_name, name, or business name)",
        )
    batch = LeadImport(
        user_id=user_id,
        filename=(filename or "leads.csv")[:255],
        status="processing",
    )
    db.add(batch)
    db.flush()
    imported = duplicates = invalid = created = 0
    errors: list[str] = []
    pending_ids: list[str] = []
    linked: list[str] = []
    for index, row in enumerate(reader, start=2):
        data = {}
        for raw_key, value in row.items():
            target = mapped_headers.get(raw_key)
            if target:
                data[target] = (value or "").strip()
        if not data.get("business_name"):
            invalid += 1
            errors.append(f"Row {index}: missing business name")
            continue
        if not data.get("primary_email") or "@" not in data.get("primary_email", ""):
            invalid += 1
            errors.append(f"Row {index}: skipped, business name and email are required")
            continue
        imported += 1
        existing = find_duplicate_lead(
            db,
            user_id,
            website_url=data.get("website_url", ""),
            phone=data.get("primary_phone", ""),
            business_name=data["business_name"],
            address=data.get("address", ""),
            email=data.get("primary_email", ""),
        )
        if existing:
            duplicates += 1
            # Already one of your leads: link it to this upload so it is part of this file and its campaign.
            if (existing.primary_email or "").strip() and existing.id not in linked:
                linked.append(existing.id)
            continue
        lead = persist_lead(db, user_id, data, source="csv", import_batch_id=batch.id)
        if not lead:
            invalid += 1
            continue
        created += 1
        if lead.website_url:
            pending_ids.append(lead.id)
    batch.lead_count = created + len(linked)
    batch.linked_lead_ids = linked
    batch.duplicate_leads = duplicates
    batch.invalid_rows = invalid
    write_audit(
        db,
        user_id=user_id,
        action="leads.imported",
        entity_type="lead_import",
        entity_id=batch.id,
        details={"imported": imported, "new": created, "duplicates": duplicates, "linked": len(linked), "invalid": invalid},
        **meta,
    )
    db.commit()
    from app.services.lead_pipeline import enqueue_lead_website_processing, refresh_import_batch

    if pending_ids:
        enqueue_lead_website_processing(pending_ids)
    refresh_import_batch(db, batch.id)
    db.refresh(batch)
    return {
        "import_id": batch.id,
        "filename": batch.filename,
        "status": batch.status,
        "imported": imported,
        "duplicates": duplicates,
        "invalid": invalid,
        "new": created,
        "queued_for_analysis": len(pending_ids),
        "errors": errors[:50],
    }


def list_imports(db: Session, user_id: str) -> list[LeadImport]:
    return (
        db.query(LeadImport)
        .filter(LeadImport.user_id == user_id)
        .order_by(LeadImport.created_at.desc())
        .limit(50)
        .all()
    )


def get_import(db: Session, user_id: str, import_id: str) -> LeadImport:
    batch = (
        db.query(LeadImport)
        .filter(LeadImport.id == import_id, LeadImport.user_id == user_id)
        .first()
    )
    if not batch:
        raise HTTPException(status_code=404, detail="Imported file not found")
    return batch


def delete_import(db: Session, user_id: str, import_id: str, meta: dict) -> None:
    batch = get_import(db, user_id, import_id)
    db.query(Lead).filter(Lead.user_id == user_id, Lead.import_batch_id == batch.id).delete(synchronize_session=False)
    db.delete(batch)
    write_audit(
        db,
        user_id=user_id,
        action="leads.import_deleted",
        entity_type="lead_import",
        entity_id=import_id,
        **meta,
    )
    db.commit()


def import_ready_lead_ids(db: Session, user_id: str, import_id: str) -> list[str]:
    batch = get_import(db, user_id, import_id)
    return [
        lead.id
        for lead in db.query(Lead)
        .filter(
            Lead.user_id == user_id,
            import_batch_filter(batch),
            Lead.status == "ready",
            Lead.primary_email != "",
            Lead.do_not_contact.is_(False),
        )
        .all()
    ]


def lead_file_info(db: Session, lead: Lead) -> dict:
    if lead.import_batch_id:
        batch = db.get(LeadImport, lead.import_batch_id)
        return {
            "file_kind": "import",
            "file_id": lead.import_batch_id,
            "file_name": batch.filename if batch else "Uploaded file",
        }
    if lead.scraping_job_id:
        from app.models import ScrapingJob

        job = db.get(ScrapingJob, lead.scraping_job_id)
        return {
            "file_kind": "scrape",
            "file_id": lead.scraping_job_id,
            "file_name": scrape_file_name(job) if job else "Maps scrape",
        }
    if lead.source == "manual":
        return {"file_kind": "manual", "file_id": "manual", "file_name": "Manual leads"}
    return {"file_kind": "", "file_id": "", "file_name": ""}


def scrape_ready_lead_ids(db: Session, user_id: str, job_id: str) -> list[str]:
    from app.models import ScrapingJob

    job = db.query(ScrapingJob).filter(ScrapingJob.id == job_id, ScrapingJob.user_id == user_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Scrape file not found")
    return [
        lead.id
        for lead in db.query(Lead)
        .filter(
            Lead.user_id == user_id,
            Lead.scraping_job_id == job.id,
            Lead.primary_email != "",
            Lead.do_not_contact.is_(False),
        )
        .all()
    ]


def scrape_file_name(job) -> str:
    if job.name and job.name not in {"Scraping job", "Maps scrape"}:
        return job.name
    parts = [f"{item.city} {item.category}".strip() for item in (job.searches or []) if item.city or item.category]
    return " + ".join(parts)[:255] or "Maps scrape"


def list_lead_files(db: Session, user_id: str) -> list[dict]:
    from app.models import ScrapingJob
    from sqlalchemy.orm import joinedload

    files: list[dict] = []
    jobs = (
        db.query(ScrapingJob)
        .options(joinedload(ScrapingJob.searches))
        .filter(ScrapingJob.user_id == user_id)
        .order_by(ScrapingJob.created_at.desc())
        .limit(50)
        .all()
    )
    for job in jobs:
        ready = (
            db.query(func.count(Lead.id))
            .filter(Lead.user_id == user_id, Lead.scraping_job_id == job.id, Lead.primary_email != "")
            .scalar()
            or 0
        )
        detail = " · ".join(
            f"{item.city}, {item.country} · {item.category}" for item in (job.searches or [])
        )
        files.append(
            {
                "id": job.id,
                "kind": "scrape",
                "name": scrape_file_name(job),
                "status": job.status,
                "lead_count": job.new_leads or ready,
                "ready_leads": ready,
                "detail": detail,
                "created_at": job.created_at,
            }
        )
    for batch in list_imports(db, user_id):
        files.append(
            {
                "id": batch.id,
                "kind": "import",
                "name": batch.filename,
                "status": batch.status,
                "lead_count": batch.lead_count,
                "ready_leads": batch.ready_leads,
                "detail": "Uploaded CSV",
                "created_at": batch.created_at,
            }
        )
    manual_ready = (
        db.query(func.count(Lead.id))
        .filter(
            Lead.user_id == user_id,
            Lead.source == "manual",
            Lead.import_batch_id.is_(None),
            Lead.scraping_job_id.is_(None),
            Lead.primary_email != "",
        )
        .scalar()
        or 0
    )
    if manual_ready:
        files.append(
            {
                "id": "manual",
                "kind": "manual",
                "name": "Manual leads",
                "status": "ready",
                "lead_count": manual_ready,
                "ready_leads": manual_ready,
                "detail": "Added one at a time",
                "created_at": None,
            }
        )
    # "Manual leads" has no date; compare it as the oldest (timezone-aware, like the other rows).
    files.sort(key=lambda row: row["created_at"] or datetime.min.replace(tzinfo=UTC), reverse=True)
    return files


def serialize_import(batch: LeadImport) -> dict:
    return {
        "id": batch.id,
        "filename": batch.filename,
        "status": batch.status,
        "lead_count": batch.lead_count,
        "ready_leads": batch.ready_leads,
        "failed_leads": batch.failed_leads,
        "analyzed_leads": batch.analyzed_leads,
        "duplicate_leads": batch.duplicate_leads,
        "invalid_rows": batch.invalid_rows,
        "error_message": batch.error_message,
        "created_at": batch.created_at,
    }


def _lead_campaign_status(db: Session, lead_id: str) -> str:
    from app.models import CampaignLead

    item = (
        db.query(CampaignLead)
        .filter(CampaignLead.lead_id == lead_id)
        .order_by(CampaignLead.updated_at.desc())
        .first()
    )
    return item.status if item else ""


def lead_detail_payload(db: Session, lead: Lead) -> dict:
    website_payload = None
    if lead.website:
        analysis = (
            db.query(WebsiteAnalysis).filter(WebsiteAnalysis.website_id == lead.website.id).first()
        )
        website_payload = {
            "id": lead.website.id,
            "url": lead.website.url,
            "domain": lead.website.domain,
            "title": lead.website.title,
            "meta_description": lead.website.meta_description,
            "https": lead.website.https,
            "reachable": lead.website.reachable,
            "platform_name": lead.website.platform_name,
            "platform_confidence": lead.website.platform_confidence,
            "last_analyzed_at": lead.website.last_analyzed_at,
            "pages_crawled": lead.website.pages_crawled,
            "social_links": lead.website.social_links,
            "analysis": None
            if not analysis
            else {
                "status": analysis.status,
                "summary": analysis.summary,
                "business_summary": analysis.business_summary,
                "business_type": analysis.business_type,
                "services": analysis.services,
                "products": analysis.products,
                "main_offering": analysis.main_offering,
                "target_audience": analysis.target_audience,
                "issues": analysis.issues,
                "opportunities": analysis.opportunities,
                "recommended_services": analysis.recommended_services,
                "category_scores": analysis.category_scores,
                "website_score": analysis.website_score,
                "analyzed_at": analysis.analyzed_at,
                "raw_data": analysis.raw_data,
            },
        }
    return {
        "id": lead.id,
        "business_name": lead.business_name,
        "contact_name": lead.contact_name,
        "category": lead.category,
        "city": lead.city,
        "country": lead.country,
        "address": lead.address,
        "primary_phone": lead.primary_phone,
        "primary_email": lead.primary_email,
        "website_url": lead.website_url,
        "website_domain": lead.website_domain,
        "platform": lead.platform,
        "website_score": lead.website_score,
        "lead_score": lead.lead_score,
        "lead_score_reasons": lead.lead_score_reasons,
        "email_status": lead.email_status,
        "email_verification_status": lead.email_verification_status,
        "scraped_at": lead.scraped_at,
        "do_not_contact": lead.do_not_contact,
        "google_maps_url": lead.google_maps_url,
        "google_place_id": lead.google_place_id,
        "rating": lead.rating,
        "review_count": lead.review_count,
        "notes": lead.notes,
        "source": lead.source,
        "status": lead.status,
        "import_batch_id": lead.import_batch_id,
        "scraping_job_id": lead.scraping_job_id,
        **lead_file_info(db, lead),
        "created_at": lead.created_at,
        "last_contacted_at": lead.last_contacted_at,
        "campaign_status": _lead_campaign_status(db, lead.id),
        "emails": lead.emails,
        "phones": lead.phones,
        "website": website_payload,
        "outreach": outreach_for_lead(db, lead),
    }
